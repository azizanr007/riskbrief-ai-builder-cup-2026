"""Build an English competition PPTX from verified facts only.

Convert the PPTX with LibreOffice to produce its matching submission PDF.
Requires python-pptx locally; never reads credentials or Firebase configuration.
"""
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'RiskBrief-AI-Builder-Cup-2026-Deck.pptx'
INK = '16382B'
MUTED = '4A6358'
GREEN = '2E7248'
CREAM = 'F5F7F2'
PALE = 'E4EEDF'
WHITE = 'FFFFFF'
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = 'RiskBrief BFSI — Evidence-Led Risk Intelligence'
prs.core_properties.subject = 'AI Builder Cup 2026 — BFSI'
prs.core_properties.author = 'RiskBrief'
prs.core_properties.comments = 'Verification snapshot: 13 October 2026. Cloud execution and participant eligibility remain unverified.'

def color(value):
    return RGBColor.from_string(value)

def text(slide, value, x, y, w, h, size=19, ink=INK, bold=False, href=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    for i, line in enumerate(value.split('\n')):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(9)
        paragraph.line_spacing = 1.1
        run = paragraph.add_run()
        run.text = line
        run.font.name = 'Liberation Sans'
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color(ink)
        if href:
            run.hyperlink.address = href
    return shape

def box(slide, x, y, w, h, fill=PALE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background()
    return shape

def slide(title, subtitle=None, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color(INK if dark else CREAM)
    fg = WHITE if dark else INK
    text(s, 'RISKBRIEF  /  AI BUILDER CUP 2026', 0.65, 0.3, 10, 0.3, 11, fg, True)
    text(s, title, 0.65, 0.95, 12, 0.8, 33, fg, True)
    if subtitle:
        text(s, subtitle, 0.65, 1.8, 12, 0.7, 18, WHITE if dark else MUTED)
    text(s, 'SYNTHETIC-DATA PROTOTYPE · NO FINANCIAL DECISIONS', 0.65, 7.08, 11.2, 0.23, 9, fg)
    text(s, str(len(prs.slides)), 12.15, 7.08, 0.45, 0.23, 9, fg)
    return s

s = slide('RiskBrief BFSI', dark=True)
text(s, 'Evidence-led risk intelligence\nwith constrained AI and human review', 0.65, 2.05, 11.8, 1.7, 37, WHITE)
text(s, 'BFSI — Intelligent Risk, Fraud & Financial Experiences', 0.65, 4.3, 11.7, 0.7, 22, 'BDDCA8')
text(s, 'Verified local implementation · 13 October 2026\nGoogle cloud execution and submission remain pending authentication.', 0.65, 5.6, 11.7, 0.9, 18, WHITE)

s = slide('A risk brief should show its evidence', 'Aggregate triage for analysts and reviewers in banking, financial services and insurance.')
box(s, 0.65, 2.7, 5.85, 3.65)
box(s, 6.82, 2.7, 5.85, 3.65)
text(s, 'The problem', 0.95, 2.97, 5.15, 0.5, 23, GREEN, True)
text(s, 'Concentration, delinquency changes and alert volumes can be difficult to scan consistently.\nAn unsupported generated summary can hide its source evidence and uncertainty.', 0.95, 3.65, 5.15, 2.2, 20)
text(s, 'The prototype solution', 7.12, 2.97, 5.15, 0.5, 23, GREEN, True)
text(s, 'Calculate illustrative indicators from synthetic data.\nAsk Gemini to select supported findings.\nValidate each citation, then let an authorized human reviewer decide.', 7.12, 3.65, 5.15, 2.2, 19)

s = slide('One inspectable workflow', 'Model-selected findings remain separate from deterministic facts and the human decision.')
steps = [
    ('01 · Submit', 'Provisioned analyst signs in and loads strict synthetic JSON.'),
    ('02 · Compute', 'Local thresholds produce evidence IDs and a reproducible input hash.'),
    ('03 · Select', 'Firebase AI Logic asks Gemini for allowed IDs and finding codes only.'),
    ('04 · Validate', 'Reject unknown citations, unsupported findings and missing uncertainty.'),
    ('05 · Review', 'Save a pending brief. An active reviewer approves or rejects once.'),
]
for i, (heading, body) in enumerate(steps):
    x = 0.65 + i * 2.46
    box(s, x, 2.85, 2.22, 3.55)
    text(s, heading, x + 0.16, 3.1, 1.94, 0.75, 21, GREEN, True)
    text(s, body, x + 0.16, 4.05, 1.94, 1.98, 17)

s = slide('Google-native web demonstration', 'The original Python/SQLite v0.4 service stays separate; Firebase Hosting serves static assets.')
stack = [
    ('Firebase Hosting', 'Existing target site configured in source; authenticated deployment remains unverified.'),
    ('Firebase Authentication', 'Sign-in-only interface for provisioned accounts; role lookup fails closed.'),
    ('Firebase AI Logic', 'Official SDK + Gemini Developer API; candidate gemini-3.5-flash-lite; no Gemini key in client.'),
    ('Cloud Firestore', 'Document ownership, active reviewer permissions, one final decision, append-only client events.'),
    ('Firebase App Check', 'Production reCAPTCHA Enterprise required; enforcement, API restrictions and quota must be verified.'),
]
for i, (name, detail) in enumerate(stack):
    y = 2.63 + i * 0.79
    box(s, 0.65, y, 3.25, 0.66)
    text(s, name, 0.84, y + 0.11, 2.9, 0.46, 17, GREEN, True)
    text(s, detail, 4.18, y + 0.04, 8.42, 0.69, 16)

s = slide('Synthetic evidence drives every finding', 'Illustrative, uncalibrated test data — no real portfolio or customer information.')
box(s, 0.65, 2.7, 6.1, 3.85)
text(s, 'E001 · Retail / synthetic example', 0.94, 2.98, 5.5, 0.55, 23, GREEN, True)
text(s, '52% exposure share\n9.1% delinquency; prior 6.1%\n3.0 percentage-point deterioration\n12 illustrative fraud alerts', 0.94, 3.75, 5.5, 2.25, 24)
text(s, 'Constrained model contract', 7.08, 2.97, 5.35, 0.55, 22, GREEN, True)
text(s, 'Input: IDs, fixed segment labels and computed codes.\nOutput: JSON findings and three mandatory uncertainty labels.\nNo tools, actions, predictions or new measurements.\nValidated IDs render fixed evidence-linked wording.', 7.08, 3.73, 5.35, 2.65, 18)

s = slide('Governance and explicit limits', 'Informational support only; the prototype cannot adjudicate credit, fraud or investments.')
cards = [
    ('Identity & review', 'Rules isolate owner records. Reviewer membership must be active. Final decisions cannot be rewritten by clients.'),
    ('Evidence validation', 'Unknown IDs, duplicate citations and unsupported findings are rejected. Synthetic, calibration and human-review uncertainty are mandatory.'),
    ('Bounded requests', '400 output tokens, 20-second SDK timeout, one in-flight UI request. Cloud quotas and App Check still require authenticated verification.'),
    ('Trust boundary', 'A static client can be modified. Rules cannot recompute analysis or authenticate model provenance. Client audit is not independently immutable.'),
]
for i, (heading, body) in enumerate(cards):
    x, y = 0.65 + (i % 2) * 6.18, 2.65 + (i // 2) * 1.86
    box(s, x, y, 5.85, 1.65)
    text(s, heading, x + 0.23, y + 0.15, 5.36, 0.4, 20, GREEN, True)
    text(s, body, x + 0.23, y + 0.66, 5.36, 0.88, 15)

s = slide('Verification is recorded, not assumed', 'Executed local release checks on 13 October 2026. Service fixtures do not represent real Gemini calls.')
for x, count, label in [(0.65, '10', 'Domain & controls'), (4.78, '6', 'Chromium / mocked SDK'), (8.91, '7', 'Actual Firestore Emulator')]:
    box(s, x, 2.7, 3.77, 1.85)
    text(s, count, x + 0.25, 2.93, 3.2, 0.83, 45, GREEN, True)
    text(s, label, x + 0.25, 3.96, 3.2, 0.36, 15)
text(s, 'Production build: passed. Desktop 1280 px / mobile 390 px: passed; no external requests or fixture code.\nApplication npm audit: 0 known vulnerabilities. Frozen Python v0.4: 77 passed, recorded on October 10.\nReal Gemini invocation: not verified. Deployment: not verified; target endpoint returned HTTP 404.', 0.65, 4.99, 12, 1.55, 18)

s = slide('Release materials and submission gates', 'No invented deployments, usage, customers, impact measurements, participants or submission receipts.')
repo = 'https://github.com/azizanr007/riskbrief-ai-builder-cup-2026'
text(s, 'Public independent source repository', 0.65, 2.67, 12, 0.45, 21, GREEN, True)
text(s, repo, 0.65, 3.23, 12, 0.6, 18, INK, href=repo)
text(s, 'Required next evidence', 0.65, 4.03, 12, 0.45, 21, GREEN, True)
text(s, 'Authenticated Spark/billing and App Check checks; one real free-tier Gemini response; HTTPS deployment; genuine public demo video; registered team/dashboard confirmation.', 0.65, 4.65, 12, 0.92, 19)
text(s, 'Official dates checked October 13: registration closed October 11; team formation through October 14; prototype deadline October 18. Eligibility: working professionals, 21+, JAPAC, 2–4 registered members. Dashboard template and team status remain unverified.', 0.65, 5.91, 12, 0.8, 15, MUTED)
prs.save(OUT)
print(OUT)
